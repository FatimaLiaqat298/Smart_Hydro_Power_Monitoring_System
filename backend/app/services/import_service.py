import csv
import io
import logging
from datetime import datetime, date
from decimal import Decimal
from typing import Any, Dict, List, Optional
from uuid import UUID
import openpyxl
from fastapi import HTTPException, UploadFile, status
from app.db.supabase import postgrest_get, postgrest_post
from app.schemas.generation import ImportSummaryResponse
from app.schemas.user import UserResponse
from app.services.powerhouse_service import list_powerhouses, seed_default_powerhouses
from app.services.generation_calculation_service import calculate_derived_generation_fields
from app.services.activity_log_service import log_activity

logger = logging.getLogger(__name__)

COLUMN_SYNONYMS = {
    "powerhouse": ["power plant", "powerhouse", "plant", "power_house", "powerhouse_name"],
    "date": ["date", "reading_date", "reading date", "timestamp"],
    "avg_load": ["avg station load (mw)", "avg station load", "avg_station_load_mw", "generation_mw", "generation (mw)"],
    "max_load": ["maximum station load (mw)", "max station load (mw)", "max_station_load_mw", "maximum load"],
    "inflow": ["actual avg cumic", "inflow_m3s", "actual_avg_inflow_cumecs", "actual inflow", "inflow"],
    "historical_hydro": ["avg historical hydrology", "avg historical hydrology (cfs)", "historical hydrology", "avg_historical_hydrology"],
    "total_units": ["total daily units (kwh)", "total_daily_units_kwh", "total units (kwh)", "total units generated"],
    "net_energy": ["daily net energy (kwh)", "daily_net_energy_kwh", "net electricity (kwh)", "net energy (kwh)"],
    "inflow_0000": ["inflow 00:00", "inflow_00:00", "00:00 inflow", "00:00 cumecs", "inflow_0000"],
    "inflow_0400": ["inflow 04:00", "inflow_04:00", "04:00 inflow", "04:00 cumecs", "inflow_0400"],
    "inflow_0800": ["inflow 08:00", "inflow_08:00", "08:00 inflow", "08:00 cumecs", "inflow_0800"],
    "inflow_1200": ["inflow 12:00", "inflow_12:00", "12:00 inflow", "12:00 cumecs", "inflow_1200"],
    "inflow_1600": ["inflow 16:00", "inflow_16:00", "16:00 inflow", "16:00 cumecs", "inflow_1600"],
    "inflow_2000": ["inflow 20:00", "inflow_20:00", "20:00 inflow", "20:00 cumecs", "inflow_2000"],
    "remarks": ["remarks/ issues", "remarks", "issues", "remarks / issues"],
    "aux_kwh": ["aux_consumption_kwh", "auxiliary consumption", "aux consumption"]
}

def resolve_column(headers: List[str], target_key: str) -> Optional[int]:
    """Finds zero-based column index matching known synonyms."""
    for idx, col in enumerate(headers):
        if col is None:
            continue
        norm = str(col).strip().lower()
        if norm in COLUMN_SYNONYMS[target_key]:
            return idx
    return None

def parse_date_value(val: Any) -> Optional[date]:
    if isinstance(val, (datetime, date)):
        return val.date() if isinstance(val, datetime) else val
    if not val:
        return None
    val_str = str(val).strip()
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(val_str.split(" ")[0], fmt.split(" ")[0]).date()
        except ValueError:
            continue
    return None

def parse_decimal_value(val: Any) -> Optional[Decimal]:
    if val is None or str(val).strip() == "":
        return None
    try:
        parsed = Decimal(str(val).replace(",", "").strip())
    except Exception as exc:
        raise ValueError(f"Invalid numeric value '{val}'") from exc
    if not parsed.is_finite() or parsed < 0:
        raise ValueError(f"Value must be a finite non-negative number: '{val}'")
    return parsed

async def process_file_import(
    file: UploadFile,
    file_type: str,
    user: UserResponse,
    client_ip: Optional[str] = None
) -> ImportSummaryResponse:
    """Processes historical data import from CSV or Excel file using pure Python libraries."""
    await seed_default_powerhouses()
    powerhouses = await list_powerhouses()
    ph_lookup = {p.name.strip().lower(): p for p in powerhouses}

    content = await file.read()
    if not content:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Uploaded file is empty.")

    raw_rows: List[List[Any]] = []
    headers: List[str] = []

    try:
        if file_type == "csv":
            text = content.decode("utf-8-sig", errors="replace")
            reader = csv.reader(io.StringIO(text))
            all_data = list(reader)
            if all_data:
                headers = all_data[0]
                raw_rows = all_data[1:]
        elif file_type == "xlsx":
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
            sheet_name = "Generation Data" if "Generation Data" in wb.sheetnames else wb.sheetnames[0]
            sheet = wb[sheet_name]
            all_data = list(sheet.iter_rows(values_only=True))
            if all_data:
                headers = [str(c) if c is not None else "" for c in all_data[0]]
                raw_rows = all_data[1:]
        else:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Unsupported file format '{file_type}'.")
    except Exception as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Error parsing {file_type} file: {str(e)}")

    col_ph = resolve_column(headers, "powerhouse")
    col_date = resolve_column(headers, "date")
    col_avg_load = resolve_column(headers, "avg_load")
    col_max_load = resolve_column(headers, "max_load")
    col_inflow = resolve_column(headers, "inflow")
    col_hist_hydro = resolve_column(headers, "historical_hydro")
    col_total_units = resolve_column(headers, "total_units")
    col_net_energy = resolve_column(headers, "net_energy")
    inflow_slot_columns = {
        slot: resolve_column(headers, f"inflow_{slot.replace(':', '')}")
        for slot in ("00:00", "04:00", "08:00", "12:00", "16:00", "20:00")
    }
    col_remarks = resolve_column(headers, "remarks")
    col_aux = resolve_column(headers, "aux_kwh")

    if col_ph is None or col_date is None or col_avg_load is None:
        missing = []
        if col_ph is None: missing.append("Power Plant")
        if col_date is None: missing.append("Date")
        if col_avg_load is None: missing.append("Avg Station Load (MW)")
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"Required columns missing in file: {', '.join(missing)}"
        )

    rows_total = len(raw_rows)
    rows_imported = 0
    rows_rejected = 0
    errors: List[Dict[str, Any]] = []
    records_to_insert: List[Dict[str, Any]] = []
    intraday_by_key: Dict[tuple, List[Dict[str, Any]]] = {}

    for idx, row in enumerate(raw_rows):
        row_num = idx + 2
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue

        raw_ph = row[col_ph] if col_ph < len(row) else None
        if raw_ph is None or not str(raw_ph).strip():
            rows_rejected += 1
            continue

        ph_name = str(raw_ph).strip()
        ph = ph_lookup.get(ph_name.lower())
        if not ph:
            rows_rejected += 1
            if len(errors) < 50:
                errors.append({"row": row_num, "error": f"Unknown powerhouse '{ph_name}'"})
            continue

        raw_date = row[col_date] if col_date < len(row) else None
        parsed_date = parse_date_value(raw_date)
        if not parsed_date:
            rows_rejected += 1
            if len(errors) < 50:
                errors.append({"row": row_num, "error": f"Invalid or missing date '{raw_date}'"})
            continue

        try:
            avg_mw = parse_decimal_value(row[col_avg_load] if col_avg_load < len(row) else None)
            if avg_mw is None:
                raise ValueError("Average station load is missing.")
        except ValueError as exc:
            rows_rejected += 1
            if len(errors) < 50:
                errors.append({"row": row_num, "error": str(exc)})
            continue

        raw_remarks = row[col_remarks] if col_remarks is not None and col_remarks < len(row) else None
        remarks_val = str(raw_remarks).strip() if raw_remarks is not None else ""

        try:
            max_mw = parse_decimal_value(row[col_max_load]) if col_max_load is not None and col_max_load < len(row) else None
            inflow = parse_decimal_value(row[col_inflow]) if col_inflow is not None and col_inflow < len(row) else None
            hist_hydro = parse_decimal_value(row[col_hist_hydro]) if col_hist_hydro is not None and col_hist_hydro < len(row) else None
            total_units = parse_decimal_value(row[col_total_units]) if col_total_units is not None and col_total_units < len(row) else None
            net_energy = parse_decimal_value(row[col_net_energy]) if col_net_energy is not None and col_net_energy < len(row) else None
            aux_kwh = parse_decimal_value(row[col_aux]) if col_aux is not None and col_aux < len(row) else None
            slot_values = []
            for slot, column_index in inflow_slot_columns.items():
                if column_index is not None and column_index < len(row):
                    slot_value = parse_decimal_value(row[column_index])
                    if slot_value is not None:
                        slot_values.append((slot, slot_value))
        except ValueError as exc:
            rows_rejected += 1
            if len(errors) < 50:
                errors.append({"row": row_num, "error": str(exc)})
            continue

        if slot_values:
            inflow = sum((value for _, value in slot_values), Decimal("0")) / Decimal(len(slot_values))

        try:
            calc = calculate_derived_generation_fields(
                avg_station_load_mw=avg_mw,
                installed_capacity_mw=ph.installed_capacity_mw,
                tariff_rate=ph.tariff_rate or Decimal("5.0000"),
                avg_annual_design_plant_factor_pct=ph.avg_annual_design_plant_factor_pct or Decimal("65.00"),
                aux_consumption_kwh=aux_kwh,
                remarks=remarks_val,
                total_daily_units_kwh=total_units,
                daily_net_energy_kwh=net_energy
            )
        except ValueError as exc:
            rows_rejected += 1
            if len(errors) < 50:
                errors.append({"row": row_num, "error": str(exc)})
            continue

        rec = {
            "powerhouse_id": str(ph.id),
            "reading_date": parsed_date.isoformat(),
            "time_slot": None,
            "max_station_load_mw": float(max_mw) if max_mw is not None else None,
            "avg_station_load_mw": float(avg_mw),
            "total_daily_units_kwh": float(calc["total_daily_units_kwh"]),
            "aux_consumption_kwh": float(calc["aux_consumption_kwh"]),
            "daily_net_energy_kwh": float(calc["daily_net_energy_kwh"]),
            "daily_plant_factor_pct": float(calc["daily_plant_factor_pct"]),
            "expected_generation_gwh": float(calc["expected_generation_gwh"]),
            "expected_revenue_pkr": float(calc["expected_revenue_pkr"]),
            "actual_revenue_pkr": float(calc["actual_revenue_pkr"]),
            "revenue_difference_pkr": float(calc["revenue_difference_pkr"]),
            "tariff_rate": float(calc["tariff_rate"]),
            "avg_annual_design_plant_factor_pct": float(calc["avg_annual_design_plant_factor_pct"]),
            "actual_avg_inflow_cumecs": float(inflow) if inflow is not None else None,
            "avg_historical_hydrology": float(hist_hydro) if hist_hydro is not None else None,
            "remarks": remarks_val if remarks_val else None,
            "outage_flag": calc["outage_flag"],
            "canal_closure_flag": calc["canal_closure_flag"],
            "entered_by": str(user.id),
            "source": f"{file_type}_import"
        }
        records_to_insert.append(rec)
        if slot_values:
            intraday_by_key[(str(ph.id), parsed_date.isoformat())] = [
                {
                    "powerhouse_id": str(ph.id),
                    "reading_date": parsed_date.isoformat(),
                    "time_slot": f"{slot}:00",
                    "station_load_mw": float(avg_mw),
                    "inflow_cumecs": float(value),
                }
                for slot, value in slot_values
            ]

    batch_size = 200
    successfully_inserted: List[Dict[str, Any]] = []
    for i in range(0, len(records_to_insert), batch_size):
        chunk = records_to_insert[i:i + batch_size]
        try:
            await postgrest_post("generation_records", chunk)
            rows_imported += len(chunk)
            successfully_inserted.extend(chunk)
        except Exception as e:
            logger.error(f"Batch import chunk failed: {e}")
            for single_rec in chunk:
                try:
                    await postgrest_post("generation_records", single_rec)
                    rows_imported += 1
                    successfully_inserted.append(single_rec)
                except Exception as ex:
                    rows_rejected += 1
                    if len(errors) < 50:
                        errors.append({"date": single_rec["reading_date"], "error": str(ex)})

    for i in range(0, len(successfully_inserted), batch_size):
        chunk = successfully_inserted[i:i + batch_size]
        slot_rows = [
            row
            for record in chunk
            for row in intraday_by_key.get((record["powerhouse_id"], record["reading_date"]), [])
        ]
        if slot_rows:
            try:
                await postgrest_post("intraday_readings", slot_rows)
            except Exception as exc:
                logger.error(f"Intraday import chunk failed: {exc}")
                if len(errors) < 50:
                    errors.append({"error": "Daily records were imported, but some intraday readings could not be saved."})

    import_audit = {
        "imported_by": str(user.id),
        "file_name": file.filename or f"import.{file_type}",
        "file_type": file_type,
        "rows_total": rows_total,
        "rows_imported": rows_imported,
        "rows_rejected": rows_rejected,
        "error_summary": errors
    }
    audit_res = await postgrest_post("data_imports", import_audit)
    import_id = UUID(audit_res[0]["id"] if isinstance(audit_res, list) else audit_res["id"])

    await log_activity(
        user_id=user.id,
        action=f"bulk_import_{file_type}",
        entity_type="data_import",
        entity_id=import_id,
        details={"file": file.filename, "imported": rows_imported, "rejected": rows_rejected},
        ip_address=client_ip
    )

    return ImportSummaryResponse(
        import_id=import_id,
        file_name=file.filename or f"import.{file_type}",
        file_type=file_type,
        rows_total=rows_total,
        rows_imported=rows_imported,
        rows_rejected=rows_rejected,
        errors=errors
    )
