import io
import logging
from datetime import date, datetime
from decimal import Decimal
from typing import Any, Dict, List, Optional
from uuid import UUID

from fastapi import HTTPException, status
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

from app.core.config import settings
from app.db.supabase import postgrest_get, postgrest_post
from app.schemas.user import UserResponse, UserRole
from app.services.activity_log_service import log_activity
from app.services.notification_service import create_notification_for_users
from app.services.email_service import send_report_ready_email

logger = logging.getLogger(__name__)


async def _calculate_report_data(
    powerhouse_id: Optional[UUID],
    date_from: date,
    date_to: date
) -> Dict[str, Any]:
    """Calculates structured report data from generation_records within the date range."""
    params: Dict[str, Any] = {
        "select": "*,powerhouses(name,installed_capacity_mw)",
        "reading_date": f"gte.{date_from.isoformat()}",
        "order": "reading_date.asc"
    }
    if powerhouse_id:
        params["powerhouse_id"] = f"eq.{str(powerhouse_id)}"
    params["reading_date"] = f"gte.{date_from.isoformat()}&reading_date=lte.{date_to.isoformat()}"
    # Use separate filters
    records_all = await postgrest_get("generation_records", {
        "select": "*,powerhouses(name,installed_capacity_mw)",
        "powerhouse_id": f"eq.{str(powerhouse_id)}" if powerhouse_id else None,
        "reading_date": f"gte.{date_from.isoformat()}",
        "order": "reading_date.asc"
    })
    # Filter to date_to in Python since PostgREST compound filters need special handling
    records = [r for r in records_all if r.get("reading_date", "") <= date_to.isoformat()]
    if powerhouse_id:
        records = [r for r in records if r.get("powerhouse_id") == str(powerhouse_id)]

    # Anomalies in period
    anomaly_params: Dict[str, Any] = {"select": "id,severity,status", "anomaly_date": f"gte.{date_from.isoformat()}"}
    if powerhouse_id:
        anomaly_params["powerhouse_id"] = f"eq.{str(powerhouse_id)}"
    anomalies_all = await postgrest_get("anomaly_alerts", anomaly_params)
    anomalies = [a for a in anomalies_all if a.get("anomaly_date", "") <= date_to.isoformat()]

    if not records:
        return {
            "period_start": date_from.isoformat(), "period_end": date_to.isoformat(),
            "powerhouse_id": str(powerhouse_id) if powerhouse_id else None,
            "total_records": 0, "total_generation_kwh": 0, "total_generation_gwh": 0,
            "avg_daily_mw": 0, "avg_inflow_cumecs": 0, "total_revenue_pkr": 0,
            "outage_days": 0, "canal_closure_days": 0, "anomaly_count": len(anomalies),
            "powerhouse_breakdown": [], "daily_data": [],
            "message": "No generation data available for the selected period."
        }

    total_kwh = sum(float(r.get("total_daily_units_kwh") or 0) for r in records)
    avg_mw = sum(float(r.get("avg_station_load_mw") or 0) for r in records) / len(records)
    inflows = [float(r["actual_avg_inflow_cumecs"]) for r in records if r.get("actual_avg_inflow_cumecs") is not None]
    total_revenue = sum(float(r.get("actual_revenue_pkr") or 0) for r in records)
    outage_days = sum(1 for r in records if r.get("outage_flag"))
    canal_days = sum(1 for r in records if r.get("canal_closure_flag"))

    # Per-powerhouse breakdown
    ph_groups: Dict[str, List] = {}
    for r in records:
        pid = r.get("powerhouse_id")
        if pid not in ph_groups:
            ph_groups[pid] = []
        ph_groups[pid].append(r)

    breakdown = []
    for pid, recs in ph_groups.items():
        ph_info = recs[0].get("powerhouses") or {}
        ph_name = ph_info.get("name") if isinstance(ph_info, dict) else "Unknown"
        ph_kwh = sum(float(r.get("total_daily_units_kwh") or 0) for r in recs)
        ph_mw = sum(float(r.get("avg_station_load_mw") or 0) for r in recs) / len(recs)
        ph_rev = sum(float(r.get("actual_revenue_pkr") or 0) for r in recs)
        ph_outage = sum(1 for r in recs if r.get("outage_flag"))
        breakdown.append({
            "powerhouse_id": pid, "powerhouse_name": ph_name,
            "total_kwh": round(ph_kwh, 2), "avg_mw": round(ph_mw, 3),
            "total_revenue_pkr": round(ph_rev, 2), "outage_days": ph_outage,
            "records": len(recs)
        })

    daily_data = [
        {
            "date": r.get("reading_date"), "avg_mw": float(r.get("avg_station_load_mw") or 0),
            "kwh": float(r.get("total_daily_units_kwh") or 0),
            "inflow": float(r.get("actual_avg_inflow_cumecs") or 0) if r.get("actual_avg_inflow_cumecs") else None,
            "revenue": float(r.get("actual_revenue_pkr") or 0),
            "outage": r.get("outage_flag"), "canal_closure": r.get("canal_closure_flag")
        }
        for r in records[:100]  # Cap daily data at 100 rows in JSON
    ]

    return {
        "period_start": date_from.isoformat(), "period_end": date_to.isoformat(),
        "powerhouse_id": str(powerhouse_id) if powerhouse_id else None,
        "total_records": len(records),
        "total_generation_kwh": round(total_kwh, 2),
        "total_generation_gwh": round(total_kwh / 1_000_000, 4),
        "avg_daily_mw": round(avg_mw, 3),
        "avg_inflow_cumecs": round(sum(inflows) / len(inflows), 3) if inflows else 0,
        "total_revenue_pkr": round(total_revenue, 2),
        "outage_days": outage_days, "canal_closure_days": canal_days,
        "anomaly_count": len(anomalies),
        "powerhouse_breakdown": breakdown,
        "daily_data": daily_data
    }


def _generate_pdf(report_data: Dict[str, Any]) -> bytes:
    """Generates a PDF report from the structured report data."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, title="SHPMFS Generation Report")
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph("Smart Hydro Power Monitoring & Forecasting System", styles["Heading1"]))
    elements.append(Paragraph(f"Generation Report: {report_data['period_start']} to {report_data['period_end']}", styles["Heading2"]))
    elements.append(Spacer(1, 12))

    summary_data = [
        ["Metric", "Value"],
        ["Total Records", str(report_data["total_records"])],
        ["Total Generation (KWh)", f"{report_data['total_generation_kwh']:,.2f}"],
        ["Total Generation (GWh)", f"{report_data['total_generation_gwh']:,.4f}"],
        ["Avg Daily Generation (MW)", f"{report_data['avg_daily_mw']:,.3f}"],
        ["Avg Inflow (Cumecs)", f"{report_data['avg_inflow_cumecs']:,.3f}"],
        ["Total Revenue (PKR)", f"{report_data['total_revenue_pkr']:,.2f}"],
        ["Outage Days", str(report_data["outage_days"])],
        ["Canal Closure Days", str(report_data["canal_closure_days"])],
        ["Anomalies", str(report_data["anomaly_count"])]
    ]
    t = Table(summary_data, colWidths=[220, 200])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e40af")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.whitesmoke, colors.white]),
    ]))
    elements.append(t)
    elements.append(Spacer(1, 20))

    if report_data.get("powerhouse_breakdown"):
        elements.append(Paragraph("Powerhouse Breakdown", styles["Heading3"]))
        bd = report_data["powerhouse_breakdown"]
        bd_data = [["Powerhouse", "Records", "Total KWh", "Avg MW", "Revenue (PKR)", "Outages"]]
        for row in bd:
            bd_data.append([
                row["powerhouse_name"], str(row["records"]),
                f"{row['total_kwh']:,.0f}", f"{row['avg_mw']:.2f}",
                f"{row['total_revenue_pkr']:,.0f}", str(row["outage_days"])
            ])
        bt = Table(bd_data, repeatRows=1)
        bt.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e40af")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.whitesmoke, colors.white]),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
        ]))
        elements.append(bt)

    doc.build(elements)
    return buf.getvalue()


def _generate_xlsx(report_data: Dict[str, Any]) -> bytes:
    """Generates an Excel report from the structured report data."""
    wb = openpyxl.Workbook()
    ws_summary = wb.active
    ws_summary.title = "Summary"
    header_fill = PatternFill(start_color="1e40af", end_color="1e40af", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    ws_summary.append(["SHPMFS Generation Report", "", ""])
    ws_summary.append([f"Period: {report_data['period_start']} to {report_data['period_end']}", "", ""])
    ws_summary.append([])
    ws_summary.append(["Metric", "Value"])
    for row in ws_summary.iter_rows(min_row=4, max_row=4):
        for cell in row:
            cell.font = header_font
            cell.fill = header_fill

    summary_rows = [
        ("Total Records", report_data["total_records"]),
        ("Total Generation (KWh)", report_data["total_generation_kwh"]),
        ("Total Generation (GWh)", report_data["total_generation_gwh"]),
        ("Avg Daily MW", report_data["avg_daily_mw"]),
        ("Avg Inflow (Cumecs)", report_data["avg_inflow_cumecs"]),
        ("Total Revenue (PKR)", report_data["total_revenue_pkr"]),
        ("Outage Days", report_data["outage_days"]),
        ("Canal Closure Days", report_data["canal_closure_days"]),
        ("Anomalies", report_data["anomaly_count"]),
    ]
    for r in summary_rows:
        ws_summary.append(list(r))

    if report_data.get("powerhouse_breakdown"):
        ws_bd = wb.create_sheet("Powerhouse Breakdown")
        headers = ["Powerhouse", "Records", "Total KWh", "Avg MW", "Revenue (PKR)", "Outage Days"]
        ws_bd.append(headers)
        for cell in ws_bd[1]:
            cell.font = header_font
            cell.fill = header_fill
        for row in report_data["powerhouse_breakdown"]:
            ws_bd.append([
                row["powerhouse_name"], row["records"], row["total_kwh"],
                row["avg_mw"], row["total_revenue_pkr"], row["outage_days"]
            ])

    if report_data.get("daily_data"):
        ws_daily = wb.create_sheet("Daily Data")
        d_headers = ["Date", "Avg MW", "Total KWh", "Inflow (Cumecs)", "Revenue (PKR)", "Outage", "Canal Closure"]
        ws_daily.append(d_headers)
        for cell in ws_daily[1]:
            cell.font = header_font
            cell.fill = header_fill
        for row in report_data["daily_data"]:
            ws_daily.append([
                row["date"], row["avg_mw"], row["kwh"],
                row.get("inflow"), row["revenue"], row["outage"], row["canal_closure"]
            ])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def generate_report(
    date_from: date,
    date_to: date,
    format_: str,
    powerhouse_id: Optional[UUID],
    user: UserResponse,
    client_ip: Optional[str] = None
) -> Dict[str, Any]:
    """Generates, stores, and returns a report for the given date range and powerhouse."""
    if date_to < date_from:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="date_to must be >= date_from")

    report_data = await _calculate_report_data(powerhouse_id, date_from, date_to)

    # Generate file bytes
    if format_ == "pdf":
        file_bytes = _generate_pdf(report_data)
        content_type = "application/pdf"
        suffix = "pdf"
    else:
        file_bytes = _generate_xlsx(report_data)
        content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        suffix = "xlsx"

    # Upload to Supabase Storage
    file_url = None
    file_path = f"reports/{date_from}_{date_to}_{datetime.utcnow().strftime('%Y%m%d%H%M%S')}.{suffix}"
    try:
        from app.db.supabase import get_supabase_admin
        supabase = get_supabase_admin()
        supabase.storage.from_(settings.SUPABASE_REPORTS_BUCKET).upload(
            path=file_path,
            file=file_bytes,
            file_options={"content-type": content_type}
        )
        file_url = supabase.storage.from_(settings.SUPABASE_REPORTS_BUCKET).get_public_url(file_path)
    except Exception as e:
        logger.warning(f"Could not upload report to Supabase Storage: {e}. Report data available in DB.")

    # Save to weekly_reports
    period_type = "custom"
    days = (date_to - date_from).days
    if days == 6 or days == 7:
        period_type = "weekly"
    elif 28 <= days <= 31:
        period_type = "monthly"

    report_record = {
        "powerhouse_id": str(powerhouse_id) if powerhouse_id else None,
        "period_type": period_type,
        "period_start": date_from.isoformat(),
        "period_end": date_to.isoformat(),
        "total_generation_mwh": report_data.get("total_generation_kwh", 0) / 1000,
        "avg_inflow_cumecs": report_data.get("avg_inflow_cumecs"),
        "total_revenue_pkr": report_data.get("total_revenue_pkr"),
        "anomaly_count": report_data.get("anomaly_count", 0),
        "summary_json": report_data,
        "file_url": file_url,
        "generated_by": str(user.id)
    }
    saved = await postgrest_post("weekly_reports", report_record)
    saved_report = saved[0] if isinstance(saved, list) else saved

    await log_activity(user.id, "generate_report", "report", UUID(saved_report["id"]),
                       {"period_start": date_from.isoformat(), "period_end": date_to.isoformat(),
                        "format": format_, "powerhouse_id": str(powerhouse_id) if powerhouse_id else "all"},
                       client_ip)

    # Notify admins
    try:
        await create_notification_for_users(
            roles=[UserRole.ADMIN],
            type_="report_ready",
            title="Report Ready",
            message=f"Report for {date_from} to {date_to} has been generated.",
            report_id=UUID(saved_report["id"])
        )
        await send_report_ready_email(saved_report["id"], date_from.isoformat(), date_to.isoformat(), file_url)
    except Exception as e:
        logger.warning(f"Report notification/email failed: {e}")

    return {**saved_report, "file_bytes": file_bytes, "content_type": content_type, "suffix": suffix}


async def list_reports(powerhouse_id: Optional[UUID] = None, page: int = 1, page_size: int = 20):
    params: Dict[str, Any] = {
        "select": "*", "order": "created_at.desc",
        "limit": str(page_size), "offset": str((page - 1) * page_size)
    }
    if powerhouse_id:
        params["powerhouse_id"] = f"eq.{str(powerhouse_id)}"
    data = await postgrest_get("weekly_reports", params)
    count_p: Dict[str, Any] = {"select": "id"}
    if powerhouse_id:
        count_p["powerhouse_id"] = f"eq.{str(powerhouse_id)}"
    total = len(await postgrest_get("weekly_reports", count_p))
    return data, total


async def get_report_by_id(report_id: UUID) -> Dict[str, Any]:
    data = await postgrest_get("weekly_reports", {"id": f"eq.{str(report_id)}"})
    if not data:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Report not found")
    return data[0]
