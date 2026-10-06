import asyncio
import logging
import os
import sys
import openpyxl
from datetime import datetime, date
from decimal import Decimal
import httpx

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from app.core.config import settings
from app.db.supabase import postgrest_get, postgrest_post, postgrest_delete
from app.services.powerhouse_service import seed_default_powerhouses, list_powerhouses
from app.services.generation_calculation_service import calculate_derived_generation_fields

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger("clean_seed")

CLEAN_EXCEL_PATH = os.path.join(os.path.dirname(__file__), "data", "Clean_Generation_Data_Merged (1).xlsx")

def parse_date(val):
    if isinstance(val, (datetime, date)):
        return val.date() if isinstance(val, datetime) else val
    if not val:
        return None
    val_str = str(val).strip()
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(val_str.split(" ")[0], fmt.split(" ")[0]).date()
        except ValueError:
            continue
    return None

async def clear_old_generation_records():
    logger.info("Clearing old generation records from Supabase...")
    headers = {
        "apikey": settings.SUPABASE_SERVICE_ROLE_KEY,
        "Authorization": f"Bearer {settings.SUPABASE_SERVICE_ROLE_KEY}",
        "Prefer": "return=minimal"
    }
    # Delete all rows where id is not null
    async with httpx.AsyncClient(timeout=30.0) as client:
        res = await client.delete(
            f"{settings.SUPABASE_URL}/rest/v1/generation_records?id=not.is.null",
            headers=headers
        )
        logger.info(f"Cleared old records, status: {res.status_code}")

async def ensure_admin_user() -> str:
    users = await postgrest_get("users", {"select": "id,role", "role": "eq.admin"})
    if users:
        return users[0]["id"]
    all_u = await postgrest_get("users", {"select": "id"})
    if all_u:
        return all_u[0]["id"]
    admin_payload = {
        "name": "System Administrator",
        "email": "admin@shpmfs.gov.pk",
        "role": "admin",
        "status": "active"
    }
    res = await postgrest_post("users", admin_payload)
    admin_id = res[0]["id"] if isinstance(res, list) else res["id"]
    return admin_id

async def run_clean_seeding():
    await clear_old_generation_records()
    admin_id = await ensure_admin_user()

    logger.info("Seeding default powerhouses...")
    await seed_default_powerhouses()
    powerhouses = await list_powerhouses()
    ph_lookup = {p.name.strip().lower(): p for p in powerhouses}

    if not os.path.exists(CLEAN_EXCEL_PATH):
        logger.error(f"Clean Excel file not found at {CLEAN_EXCEL_PATH}")
        return

    logger.info(f"Loading clean dataset: {CLEAN_EXCEL_PATH}...")
    wb = openpyxl.load_workbook(CLEAN_EXCEL_PATH, data_only=True)
    sheet_name = "Clean_Generation_Data" if "Clean_Generation_Data" in wb.sheetnames else wb.sheetnames[0]
    sheet = wb[sheet_name]

    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        logger.error("No rows found in clean Excel sheet.")
        return

    headers = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
    logger.info(f"Clean Headers: {headers}")

    col_date = headers.index("date") if "date" in headers else 0
    col_ph = headers.index("power plant") if "power plant" in headers else 4
    col_max_load = headers.index("maximum station load (mw)") if "maximum station load (mw)" in headers else 6
    col_avg_load = headers.index("avg station load (mw)") if "avg station load (mw)" in headers else 7
    col_inflow = headers.index("actual avg cumic") if "actual avg cumic" in headers else 11
    col_hist_hydro = headers.index("avg historical hydrology") if "avg historical hydrology" in headers else 10
    col_remarks = headers.index("remarks/ issues") if "remarks/ issues" in headers else 21
    col_aux = headers.index("aux. cons.") if "aux. cons." in headers else 20

    data_rows = rows[1:]
    logger.info(f"Processing {len(data_rows)} rows from sheet '{sheet_name}'...")

    records_to_insert = []
    skipped = 0

    for r in data_rows:
        if not r or r[col_ph] is None or not str(r[col_ph]).strip():
            skipped += 1
            continue

        ph_name = str(r[col_ph]).strip()
        ph = ph_lookup.get(ph_name.lower())
        if not ph:
            skipped += 1
            continue

        parsed_date = parse_date(r[col_date])
        if not parsed_date:
            skipped += 1
            continue

        raw_avg = r[col_avg_load]
        try:
            avg_mw = Decimal(str(raw_avg)) if raw_avg is not None else Decimal("0")
        except Exception:
            avg_mw = Decimal("0")

        raw_max = r[col_max_load]
        max_mw = float(raw_max) if raw_max is not None and str(raw_max).strip() != "" else None

        raw_inflow = r[col_inflow]
        inflow = float(raw_inflow) if raw_inflow is not None and str(raw_inflow).strip() != "" else None

        raw_hist = r[col_hist_hydro]
        hist_hydro = float(raw_hist) if raw_hist is not None and str(raw_hist).strip() != "" else None

        raw_aux = r[col_aux] if col_aux < len(r) else None
        aux_kwh = Decimal(str(raw_aux)) if raw_aux is not None and str(raw_aux).strip() != "" else None

        raw_remarks = r[col_remarks] if col_remarks < len(r) else None
        remarks_val = str(raw_remarks).strip() if raw_remarks is not None and str(raw_remarks).strip() != "" else None

        calc = calculate_derived_generation_fields(
            avg_station_load_mw=avg_mw,
            installed_capacity_mw=ph.installed_capacity_mw,
            tariff_rate=ph.tariff_rate or Decimal("5.0000"),
            avg_annual_design_plant_factor_pct=ph.avg_annual_design_plant_factor_pct or Decimal("65.00"),
            aux_consumption_kwh=aux_kwh,
            remarks=remarks_val or ""
        )

        rec = {
            "powerhouse_id": str(ph.id),
            "reading_date": parsed_date.isoformat(),
            "time_slot": None,
            "max_station_load_mw": max_mw,
            "avg_station_load_mw": float(avg_mw),
            "total_daily_units_kwh": float(calc["total_daily_units_kwh"]),
            "aux_consumption_kwh": float(calc["aux_consumption_kwh"]),
            "daily_net_energy_kwh": float(calc["daily_net_energy_kwh"]),
            "daily_plant_factor_pct": float(calc["daily_plant_factor_pct"]),
            "expected_generation_gwh": float(calc["expected_generation_gwh"]),
            "expected_revenue_pkr": float(calc["expected_revenue_pkr"]),
            "actual_revenue_pkr": float(calc["actual_revenue_pkr"]),
            "revenue_difference_pkr": float(calc["revenue_difference_pkr"]),
            "tariff_rate": float(calc["tariff_rate"]),
            "avg_annual_design_plant_factor_pct": float(calc["avg_annual_design_plant_factor_pct"]),
            "actual_avg_inflow_cumecs": inflow,
            "avg_historical_hydrology": hist_hydro,
            "remarks": remarks_val,
            "outage_flag": calc["outage_flag"],
            "canal_closure_flag": calc["canal_closure_flag"],
            "entered_by": admin_id,
            "source": "clean_excel_seed"
        }
        records_to_insert.append(rec)

    logger.info(f"Prepared {len(records_to_insert)} clean records for insertion (skipped {skipped}).")

    batch_size = 200
    inserted = 0
    for i in range(0, len(records_to_insert), batch_size):
        chunk = records_to_insert[i:i + batch_size]
        try:
            await postgrest_post("generation_records", chunk)
            inserted += len(chunk)
            logger.info(f"Inserted clean batch {i//batch_size + 1}: {inserted}/{len(records_to_insert)} records...")
        except Exception as e:
            logger.warning(f"Batch insert error at chunk {i}: {e}. Retrying single rows...")
            for single_rec in chunk:
                try:
                    await postgrest_post("generation_records", single_rec)
                    inserted += 1
                except Exception:
                    pass

    logger.info(f"✅ Seeding complete! Successfully inserted {inserted} clean generation records into Supabase.")

if __name__ == "__main__":
    asyncio.run(run_clean_seeding())
